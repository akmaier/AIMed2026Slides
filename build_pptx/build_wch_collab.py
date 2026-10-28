#!/usr/bin/env python3
"""Build West_China_Collaboration_2026.pptx from output_tex/main.pdf.

The deck is typeset in LaTeX (output_tex/main.tex); PowerPoint is only the
container for presenting and for the movies.  Every PDF page becomes one
full-bleed picture on a 16:9 slide (33.867 x 19.05 cm, the size of the beamer
page).  On the Motion Compensation slide the two animated GIFs are placed on
top of their first-frame stills, at the positions main.tex fixes for them.

Run from the repository root:
  cd output_tex && xelatex main.tex && xelatex main.tex && cd ..
  python3 build_pptx/build_wch_collab.py
"""
from __future__ import annotations

import subprocess
from pathlib import Path

from pptx import Presentation
from pptx.util import Cm, Emu

ROOT = Path(__file__).resolve().parent.parent
PDF = ROOT / "output_tex" / "main.pdf"
FIG = ROOT / "output_tex" / "figures" / "wchc"
PAGES = ROOT / "build_pptx" / "pages"          # gitignored
OUT = ROOT / "build_pptx" / "West_China_Collaboration_2026.pptx"

# page -> [(gif, x_cm, y_cm, side_cm)]; must match the Motion Compensation frame in main.tex
MOTION_GIFS = {
    12: [("motion_ecg_gated.gif", 5.9, 5.1, 10.8),
         ("motion_compensated.gif", 17.17, 5.1, 10.8)],
}


def main() -> None:
    PAGES.mkdir(parents=True, exist_ok=True)
    for old in PAGES.glob("page-*.png"):
        old.unlink()
    subprocess.run(["gs", "-q", "-dNOPAUSE", "-dBATCH", "-sDEVICE=png16m", "-r200",
                    f"-sOutputFile={PAGES}/page-%02d.png", str(PDF)], check=True)
    pages = sorted(PAGES.glob("page-*.png"))

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    blank = prs.slide_layouts[6]
    for i, page in enumerate(pages, start=1):
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(page), 0, 0, prs.slide_width, prs.slide_height)
        for gif, x, y, side in MOTION_GIFS.get(i, []):
            slide.shapes.add_picture(str(FIG / gif), Cm(x), Cm(y), Cm(side), Cm(side))
    prs.save(str(OUT))
    print(f"wrote {OUT} ({len(pages)} slides)")


if __name__ == "__main__":
    main()
